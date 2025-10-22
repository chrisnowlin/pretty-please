import os
import shutil
from pathlib import Path

import pytest

from src.jina_rag_pipeline.ingestion.deepseek_loader import DeepseekLoader
from src.jina_rag_pipeline.ingestion.ocr_config import OCRConfig

RUN_DEEPSEEK = os.getenv("RUN_DEEPSEEK_INTEGRATION", "0") == "1"

pytestmark = [
    pytest.mark.integration,
    pytest.mark.slow,
    pytest.mark.skipif(
        not RUN_DEEPSEEK,
        reason="Set RUN_DEEPSEEK_INTEGRATION=1 to enable Deepseek OCR integration tests",
    ),
]


@pytest.fixture(scope="session")
def deepseek_test_env():
    """Ensure deterministic processing profile and clean checkpoints."""
    old_profile = os.environ.get("PROCESSING_PROFILE")
    os.environ["PROCESSING_PROFILE"] = "conservative"
    yield
    if old_profile is None:
        os.environ.pop("PROCESSING_PROFILE", None)
    else:
        os.environ["PROCESSING_PROFILE"] = old_profile
    shutil.rmtree("checkpoints", ignore_errors=True)


@pytest.fixture(scope="session")
def deepseek_loader(deepseek_test_env):
    """Create a reusable Deepseek loader with tiny preset for faster tests."""
    config = OCRConfig.deepseek_tiny()
    config.batch_size = 1
    config.render_workers = 1
    config.analysis_workers = 1
    config.pre_render_batches = 1
    loader = DeepseekLoader(config=config)
    yield loader


def _create_sample_pdf(tmp_path: Path) -> Path:
    """Generate a single-page PDF with simple text content."""
    import fitz  # PyMuPDF

    pdf_path = tmp_path / "deepseek_sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    rect = fitz.Rect(72, 72, 468, 300)
    page.insert_textbox(
        rect,
        "Deepseek integration test.\nThis PDF validates end-to-end OCR extraction.",
        fontsize=16,
    )
    doc.save(pdf_path)
    doc.close()
    return pdf_path


def _create_sample_pptx(tmp_path: Path) -> Path:
    """Generate a one-slide PPTX with title and body text."""
    from pptx import Presentation

    ppt_path = tmp_path / "deepseek_sample.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Deepseek OCR Integration"
    slide.placeholders[1].text = (
        "This slide ensures Deepseek can process rendered presentation content."
    )
    prs.save(ppt_path)
    return ppt_path


def _assert_document_basics(document, loader_config):
    metadata = document.metadata
    assert metadata.get("extraction_method") == "deepseek-ocr"
    assert metadata.get("ocr_engine") == "deepseek"
    assert metadata.get("resolution_mode") == loader_config.resolution_mode
    assert metadata.get("compression_enabled") == loader_config.enable_compression
    assert document.content and document.content.strip()

    regions = metadata.get("regions", [])
    assert regions, "Expected semantic regions from Deepseek output"
    assert any(getattr(region, "content", "").strip() for region in regions)
    return metadata


def test_deepseek_pdf_end_to_end(deepseek_loader: DeepseekLoader, tmp_path):
    pdf_path = _create_sample_pdf(tmp_path)
    document = deepseek_loader.load(pdf_path)
    metadata = _assert_document_basics(document, deepseek_loader.config)

    assert metadata.get("format") == "pdf"
    assert metadata.get("page_count") == 1
    assert "integration test" in document.content.lower()


def test_deepseek_pptx_end_to_end(deepseek_loader: DeepseekLoader, tmp_path):
    pptx_path = _create_sample_pptx(tmp_path)
    document = deepseek_loader.load(pptx_path)
    metadata = _assert_document_basics(document, deepseek_loader.config)

    assert metadata.get("format") in {"powerpoint", "pptx"}
    assert metadata.get("slide_count") == 1
    assert "integration" in document.content.lower()
